import csv
import json
import re
import uuid
from io import StringIO
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.config import get_settings
from app.models import Document


class DocumentInputError(ValueError):
    pass


class DocumentExtractionError(ValueError):
    pass


SAFE_FILENAME_PATTERN = re.compile(r"[^A-Za-z0-9._-]+")
TEXT_SUFFIXES = {".txt", ".md", ".json", ".csv"}
SPREADSHEET_SUFFIXES = {".xlsx"}
PDF_SUFFIXES = {".pdf"}


def _safe_filename(filename: str) -> str:
    name = Path(filename or "document").name.strip() or "document"
    safe = SAFE_FILENAME_PATTERN.sub("_", name)
    return safe[:180] or "document"


def _storage_path(original_filename: str) -> Path:
    settings = get_settings()
    upload_dir = Path(settings.upload_dir).resolve()
    upload_dir.mkdir(parents=True, exist_ok=True)
    safe_name = _safe_filename(original_filename)
    path = (upload_dir / f"{uuid.uuid4()}_{safe_name}").resolve()
    if upload_dir not in path.parents:
        raise DocumentInputError("Invalid upload path")
    return path


def save_uploaded_document(
    db: Session,
    *,
    content: bytes,
    original_filename: str,
    file_type: str,
    plant_id: str | None = None,
    scenario_id: str | None = None,
    document_category: str | None = None,
) -> Document:
    if not content:
        raise DocumentInputError("Uploaded file is empty")
    path = _storage_path(original_filename)
    path.write_bytes(content)
    record = Document(
        plant_id=plant_id or None,
        scenario_id=scenario_id or None,
        filename=path.name,
        original_filename=_safe_filename(original_filename),
        file_type=file_type or "application/octet-stream",
        storage_path=str(path),
        document_category=document_category or None,
        upload_status="uploaded",
        extraction_status="pending",
        file_size_bytes=len(content),
    )
    db.add(record)
    db.commit()
    db.refresh(record)
    return record


def list_documents(
    db: Session,
    *,
    plant_id: str | None = None,
    scenario_id: str | None = None,
) -> list[Document]:
    query = select(Document)
    if plant_id:
        query = query.where(Document.plant_id == plant_id)
    if scenario_id:
        query = query.where(Document.scenario_id == scenario_id)
    return list(db.scalars(query.order_by(Document.created_at.desc(), Document.id.desc())))


def get_document_or_raise(db: Session, document_id: str) -> Document:
    record = db.get(Document, document_id)
    if record is None:
        raise DocumentInputError("Document not found")
    return record


def _extract_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    return "\n\n".join(page.extract_text() or "" for page in reader.pages).strip()


def _extract_xlsx(path: Path) -> str:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sections: list[str] = []
    for sheet in workbook.worksheets:
        sections.append(f"Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            values = [str(value) for value in row if value is not None]
            if values:
                sections.append(" | ".join(values))
    workbook.close()
    return "\n".join(sections).strip()


def _extract_csv(path: Path) -> str:
    text = path.read_text(encoding="utf-8", errors="replace")
    reader = csv.reader(StringIO(text))
    return "\n".join(" | ".join(row) for row in reader).strip()


def _extract_text_like(path: Path, suffix: str) -> str:
    text = path.read_text(encoding="utf-8", errors="replace")
    if suffix == ".json":
        try:
            parsed = json.loads(text)
            return json.dumps(parsed, indent=2, sort_keys=True, ensure_ascii=False)
        except json.JSONDecodeError:
            return text
    return text


def extract_text_from_path(path: Path, original_filename: str) -> str:
    suffix = Path(original_filename).suffix.lower()
    if suffix in PDF_SUFFIXES:
        text = _extract_pdf(path)
    elif suffix in SPREADSHEET_SUFFIXES:
        text = _extract_xlsx(path)
    elif suffix == ".csv":
        text = _extract_csv(path)
    elif suffix in TEXT_SUFFIXES:
        text = _extract_text_like(path, suffix)
    else:
        raise DocumentExtractionError(f"Unsupported document type: {suffix or 'unknown'}")
    if not text.strip():
        raise DocumentExtractionError("No extractable text found")
    return text.strip()


def extract_document_text(db: Session, document_id: str) -> Document:
    record = get_document_or_raise(db, document_id)
    try:
        path = Path(record.storage_path).resolve()
        if not path.exists():
            raise DocumentExtractionError("Stored file was not found")
        record.extracted_text = extract_text_from_path(path, record.original_filename)
        record.extraction_status = "extracted"
        record.extraction_error = None
    except DocumentExtractionError as exc:
        record.extracted_text = None
        record.extraction_status = "failed"
        record.extraction_error = str(exc)
    db.commit()
    db.refresh(record)
    return record
